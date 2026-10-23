import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F4F5F7")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        b = OxmlElement(f'w:{border_name}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), 'D0D7DE')
        tcBorders.append(b)
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(36, 41, 47)
    doc.add_paragraph()

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    run = h.runs[0]
    if level == 1:
        run.font.name = 'Arial'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(27, 54, 93)
    elif level == 2:
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(41, 105, 176)
    elif level == 3:
        run.font.name = 'Arial'
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(51, 51, 51)
    return h

def build_master_doc():
    doc = docx.Document()
    
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)
        
    # Title Header Block
    header_table = doc.add_table(rows=1, cols=1)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = header_table.cell(0, 0)
    set_cell_background(c, "1B365D")
    set_cell_margins(c, top=200, bottom=200, left=250, right=250)
    
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("GROWTHOS: MASTER TECHNICAL SPECIFICATION & PROJECT REPORT\nInstitutional-Grade Intelligence for D2C E-Commerce Founders")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    
    p2 = c.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Course: Full Stack Development - I (PS-FSD) | Sem V | A.Y. 2026-27\n7 AI Agents · Next.js 16 · PostgreSQL · Prisma ORM · Groq LPU Inference")
    r2.font.name = 'Arial'
    r2.font.size = Pt(9.5)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(200, 220, 245)
    
    doc.add_paragraph()

    # Section 1
    add_styled_heading(doc, "1. Executive Summary & Problem Statement", level=1)
    p = doc.add_paragraph(
        "GrowthOS is a full-stack AI Growth Operating System engineered to solve the operational and margin crisis faced by "
        "direct-to-consumer (D2C) e-commerce brands. Founders currently wrestle with fragmented dashboards across Shopify, Meta Ads, "
        "Google Analytics, Shiprocket, and payment gateways. While top-line revenue may increase, hidden margin leaks such as high "
        "Customer Acquisition Costs (CAC), Cash on Delivery (COD) friction, and 18-30% Return-to-Origin (RTO) rates erode net profitability.\n\n"
        "GrowthOS unifies these metrics into a single pane of glass and triggers a parallel 7-agent AI diagnostic engine powered by Groq "
        "LPUs in under 4.0 seconds, yielding prioritized action items with step-by-step SOPs, what-if strategy simulations, and board-ready PDF reports."
    )
    p.paragraph_format.line_spacing = 1.15

    # Section 2
    add_styled_heading(doc, "2. High-Level Architecture & Tech Stack", level=1)
    stack_data = [
        ["Layer", "Technology", "Role & Engineering Rationale"],
        ["Frontend & SSR", "Next.js 16.2.12 (Turbopack) + React 19", "App Router architecture with Server Components and sub-second page transitions."],
        ["Styling System", "Tailwind CSS 3.4", "Curated dark mode palette (#0D0D0E, zinc-900) and responsive grid layout."],
        ["Database", "PostgreSQL (Supabase)", "Strict ACID compliance, relational integrity, row-level locking, and native JSONB indexing."],
        ["ORM", "Prisma 6.5.0", "Type-safe database client with connection pooling (port 6543) and schema migration management."],
        ["Authentication", "Auth.js v5 (NextAuth beta 25)", "JWT session handling with Bcrypt password hashing and Google OAuth integration."],
        ["AI Inference Engine", "Groq SDK (Llama 3.3-70B Versatile)", "Ultra-low latency LPU hardware executing 7 parallel agents in < 4.0 seconds."],
        ["Data Visualizations", "Recharts 3.10.1", "Dynamic SVG metric charts, historical trend lines, and health score gauges."],
        ["PDF Generation", "jsPDF 2.5.2 + html2canvas 1.4.1", "Pixel-perfect client-side compilation of executive reports into downloadable PDF artifacts."]
    ]
    t = doc.add_table(rows=len(stack_data), cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(stack_data):
        for c_idx, val in enumerate(row):
            cell = t.cell(r_idx, c_idx)
            set_cell_margins(cell, top=45, bottom=45, left=70, right=70)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Arial'
            if r_idx == 0:
                set_cell_background(cell, "2969B0")
                run.font.size = Pt(9.0)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "F8F9FA" if r_idx % 2 == 1 else "FFFFFF")
                run.font.size = Pt(8.0)
                if c_idx == 0:
                    run.font.bold = True
    doc.add_paragraph()

    # Section 3
    add_styled_heading(doc, "3. The 7-Agent AI Diagnostic Engine", level=1)
    p = doc.add_paragraph(
        "GrowthOS orchestrates 7 specialized AI agents concurrently via Promise.all() in src/lib/ai/orchestrator.ts. "
        "Each agent evaluates a specific operational domain, identifies bottlenecks, computes an impact estimate, and provides an actionable SOP:"
    )
    p.paragraph_format.line_spacing = 1.15

    agent_data = [
        ["#", "Agent Name", "Focus & Metrics Analyzed", "Concrete Output Example"],
        ["1", "Finance Agent", "Gross Margin, Operating Expenses, COGS, Net Margin, Pricing Tier", "Renegotiate fabric supplier terms for top 3 SKUs; recover 5% COGS margin (+₹92,000/mo)."],
        ["2", "Marketing Agent", "Blended CAC, Meta vs. Google ROAS, Ad Spend Ratio, Channel Mix", "Shift 40% of ad spend to Creator UGC video ads; recover ROAS from 2.2x to 3.1x."],
        ["3", "Operations Agent", "COD Rate %, RTO %, Fulfillment SLA, Courier NDR %", "Deploy automated WhatsApp NDR verification; block high-RTO (>35%) pincodes from COD."],
        ["4", "Technology Agent", "Cart Abandonment Rate, Mobile Drop-off %, Checkout Latency", "Install Razorpay Magic 1-Click Checkout; reduce mobile checkout drop-offs (+₹1.84L/mo)."],
        ["5", "Unit Economics Agent", "AOV, Contribution Margin 1 & 2, LTV:CAC Ratio", "Bundle low-margin products into 2-packs at ₹2,299 threshold; lift CM2 to 21.5%."],
        ["6", "Inventory Agent", "Turnover Velocity, Deadstock Capital, Stockout Rate", "Launch flash clearance for >60-day aged stock; free up ₹2.8L trapped working capital."],
        ["7", "Retention & LTV Agent", "Repeat Customer %, 30/60/90-Day Repurchase, WhatsApp Open Rate", "Launch automated WhatsApp 3-step post-delivery flow; lift 60-day repeat rate from 18% to 26%."]
    ]
    t = doc.add_table(rows=len(agent_data), cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(agent_data):
        for c_idx, val in enumerate(row):
            cell = t.cell(r_idx, c_idx)
            set_cell_margins(cell, top=45, bottom=45, left=70, right=70)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Arial'
            if r_idx == 0:
                set_cell_background(cell, "2969B0")
                run.font.size = Pt(9.0)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "F8F9FA" if r_idx % 2 == 1 else "FFFFFF")
                run.font.size = Pt(8.0)
                if c_idx in [0, 1]:
                    run.font.bold = True
    doc.add_paragraph()

    # Section 4
    add_styled_heading(doc, "4. Database Architecture & Schema Design", level=1)
    p = doc.add_paragraph(
        "The database is designed with 10 tables and 6 enums adhering to 3rd Normal Form (3NF) for relational entities, "
        "while leveraging PostgreSQL JSONB columns for dynamic multi-agent diagnostic outputs and time-series snapshots."
    )
    p.paragraph_format.line_spacing = 1.15

    schema_summary = [
        ["Entity Name", "Primary Key", "Foreign Keys", "Core Responsibilities"],
        ["User", "id (CUID)", "None", "Identity, authentication, bcrypt password hash, account timestamps."],
        ["Business", "id (CUID)", "userId -> User(id)", "Master store profile, revenue, CAC, AOV, ROAS, COD %, RTO %, health score."],
        ["BusinessSnapshot", "id (CUID)", "businessId -> Business(id)", "Immutable time-series records for longitudinal trend tracking."],
        ["Analysis", "id (CUID)", "businessId, createdById", "JSONB storage for 7-agent diagnostic results, health score, and executive summary."],
        ["Simulation", "id (CUID)", "businessId, createdById", "Hypothetical strategy runs, short/long-term forecasts, risks, and opportunities."],
        ["Report", "id (CUID)", "businessId, analysisId, simulationId", "Executive briefing documents, Chief Risk, Top Opportunity, and action plans."],
        ["Account / Session", "id (CUID)", "userId -> User(id)", "NextAuth session management and OAuth tokens with cascade deletion."]
    ]
    t = doc.add_table(rows=len(schema_summary), cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(schema_summary):
        for c_idx, val in enumerate(row):
            cell = t.cell(r_idx, c_idx)
            set_cell_margins(cell, top=45, bottom=45, left=70, right=70)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Arial'
            if r_idx == 0:
                set_cell_background(cell, "2969B0")
                run.font.size = Pt(9.0)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "F8F9FA" if r_idx % 2 == 1 else "FFFFFF")
                run.font.size = Pt(8.0)
                if c_idx in [0, 1]:
                    run.font.bold = True
    doc.add_paragraph()

    # Section 5
    add_styled_heading(doc, "5. Complete Route & Application Feature Catalog", level=1)
    routes_data = [
        ["Route", "Type", "Access", "Key Features & Functionality"],
        ["/", "Static", "Public", "Marketing landing page with value prop, 7-agent cards, and interactive CTA."],
        ["/demo", "Static", "Public", "Instant zero-config demo with pre-loaded ethnic wear D2C case study."],
        ["/login", "Static", "Public", "Email/Password and Google OAuth authentication portal."],
        ["/register", "Static", "Public", "New user registration with bcrypt password hashing."],
        ["/onboarding", "Dynamic", "Protected", "3-step wizard for configuring store metrics, logistics, and financials."],
        ["/dashboard", "Dynamic", "Protected", "Command center with composite Health Score radial gauge and metric alerts."],
        ["/dashboard/diagnose", "Dynamic", "Protected", "Interactive 7-agent diagnosis studio with SOP checklists and tool paths."],
        ["/dashboard/simulate", "Dynamic", "Protected", "What-if strategy simulation testing natural language hypotheses."],
        ["/dashboard/report", "Dynamic", "Protected", "CEO Executive Briefing with 1-click high-resolution PDF download."],
        ["/dashboard/history", "Dynamic", "Protected", "Chronological snapshot timeline and historical charts powered by Recharts."],
        ["/dashboard/settings", "Dynamic", "Protected", "Store parameters management and metric re-tuning."]
    ]
    t = doc.add_table(rows=len(routes_data), cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(routes_data):
        for c_idx, val in enumerate(row):
            cell = t.cell(r_idx, c_idx)
            set_cell_margins(cell, top=45, bottom=45, left=70, right=70)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Arial'
            if r_idx == 0:
                set_cell_background(cell, "2969B0")
                run.font.size = Pt(9.0)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "F8F9FA" if r_idx % 2 == 1 else "FFFFFF")
                run.font.size = Pt(8.0)
                if c_idx == 0:
                    run.font.bold = True
    doc.add_paragraph()

    # Section 6
    add_styled_heading(doc, "6. Presentation & Viva Defense Quick Guide", level=1)
    p = doc.add_paragraph(
        "Key questions and answers for tomorrow's project presentation:\n\n"
        "Q: Why Next.js + PostgreSQL instead of traditional Spring Boot?\n"
        "A: Next.js App Router provides modern full-stack capabilities with serverless API Route Handlers acting as REST controllers, "
        "and Prisma ORM managing PostgreSQL. This allows direct client-side PDF generation, sub-second Turbopack compilation, and easy deployment.\n\n"
        "Q: How do the 7 agents run concurrently?\n"
        "A: In src/lib/ai/orchestrator.ts, all 7 agent prompts are executed in parallel via Promise.all() against Groq's high-speed LPU inference engine, "
        "completing in under 4 seconds rather than 28 seconds sequentially.\n\n"
        "Q: Why use JSONB in PostgreSQL?\n"
        "A: JSONB gives us the best of both worlds: strict relational integrity (foreign keys, cascade deletes) for store entities, combined "
        "with flexible, schema-free document storage for the deeply nested 7-agent diagnostic outputs without needing 15+ sub-tables."
    )
    p.paragraph_format.line_spacing = 1.15

    output_path = "GROWTHOS_MASTER_DOCUMENTATION.docx"
    doc.save(output_path)
    output_path_sub = "submissions/GROWTHOS_MASTER_DOCUMENTATION.docx"
    doc.save(output_path_sub)
    print(f"Successfully generated: {output_path} and {output_path_sub}")

build_master_doc()
