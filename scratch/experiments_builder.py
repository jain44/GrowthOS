import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_image_placeholder(doc, figure_num, caption):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=180, bottom=180, left=220, right=220)
    
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        b = OxmlElement(f'w:{border_name}')
        b.set(qn('w:val'), 'dashed')
        b.set(qn('w:sz'), '6')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), '306090')
        tcBorders.append(b)
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    run1 = p.add_run("📸 [INSERT / PASTE SCREENSHOT HERE]\n")
    run1.font.name = 'Arial'
    run1.font.size = Pt(10)
    run1.font.bold = True
    run1.font.color.rgb = RGBColor(27, 54, 93)
    
    run2 = p.add_run(f"Figure {figure_num}: {caption}\n(Recommended Dimensions: 1200 x 650 px | Resolution: 300 DPI)")
    run2.font.name = 'Arial'
    run2.font.size = Pt(8.5)
    run2.font.italic = True
    run2.font.color.rgb = RGBColor(70, 80, 95)
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(2)
    p_spacer.paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
    
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        b = OxmlElement(f'w:{border_name}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), 'D0D7DE')
        tcBorders.append(b)
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(36, 41, 47)
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(2)
    p_spacer.paragraph_format.space_after = Pt(4)

def add_table_data(doc, table_rows, col_widths=None):
    t = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(table_rows):
        for c_idx, val in enumerate(row):
            cell = t.cell(r_idx, c_idx)
            set_cell_margins(cell, top=45, bottom=45, left=65, right=65)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Arial'
            if r_idx == 0:
                set_cell_background(cell, "1B365D")
                run.font.size = Pt(8.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "F4F6F9" if r_idx % 2 == 1 else "FFFFFF")
                run.font.size = Pt(8.0)
                if c_idx == 0:
                    run.font.bold = True
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(2)
    p_spacer.paragraph_format.space_after = Pt(4)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    run = h.runs[0]
    if level == 1:
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        run.font.name = 'Arial'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(27, 54, 93)
    elif level == 2:
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(41, 105, 176)
    elif level == 3:
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
        run.font.name = 'Arial'
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 41, 59)
    return h

def add_body_paragraph(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run = p.runs[0]
    run.font.name = 'Arial'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(40, 40, 40)
    return p

print("Helper functions ready.")
