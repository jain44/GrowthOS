import sys
import os

current_dir = os.path.dirname(os.path.abspath(__file__))
project_root = os.path.abspath(os.path.join(current_dir, ".."))
os.chdir(project_root)

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from experiments_common import (
    set_cell_background, set_cell_margins, add_image_placeholder,
    add_code_block, add_table_data, add_styled_heading, add_body_paragraph
)

from exp_data_part1 import EXP_1
from exp_data_part2 import EXP_2, EXP_3, EXP_4, EXP_5
from exp_data_part3 import EXP_6, EXP_7, EXP_8, EXP_9, EXP_10

ALL_EXPERIMENTS = [EXP_1, EXP_2, EXP_3, EXP_4, EXP_5, EXP_6, EXP_7, EXP_8, EXP_9, EXP_10]

def render_experiment_docx(doc, exp):
    add_styled_heading(doc, f"Experiment {exp['num']}: {exp['title']}", level=1)
    
    # 1. Aim
    add_styled_heading(doc, "1. Aim", level=2)
    add_body_paragraph(doc, exp['aim'])
    
    # 2. Objective
    add_styled_heading(doc, "2. Objective", level=2)
    for obj in exp['objective']:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(obj)
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(40, 40, 40)
    doc.add_paragraph()
    
    # 3. Theory
    add_styled_heading(doc, "3. Theory", level=2)
    for para in exp['theory_text'].strip().split('\n\n'):
        if para.strip():
            add_body_paragraph(doc, para.strip())
            
    if "comparison_table" in exp:
        add_styled_heading(doc, "Architectural Comparison & Evaluation Matrix", level=3)
        add_table_data(doc, exp['comparison_table'])
        
    # 4. Procedure
    add_styled_heading(doc, "4. Procedure", level=2)
    for para in exp['procedure_steps'].strip().split('\n\n'):
        if para.strip():
            add_body_paragraph(doc, para.strip())
            
    add_styled_heading(doc, "4.1 Maven Build Configuration (pom.xml)", level=3)
    add_code_block(doc, exp['code_pom'])
    
    if "code_props" in exp:
        add_styled_heading(doc, "4.2 Application Configuration (application.properties)", level=3)
        add_code_block(doc, exp['code_props'])
        
    if "code_main" in exp:
        add_styled_heading(doc, "4.3 Java Source Code Implementation", level=3)
        add_code_block(doc, exp['code_main'])
        
    # 5. Result and Discussion
    add_styled_heading(doc, "5. Result and Discussion", level=2)
    for para in exp['result_discussion_text'].strip().split('\n\n'):
        if para.strip():
            add_body_paragraph(doc, para.strip())
            
    if "test_matrix" in exp:
        add_styled_heading(doc, "5.1 Test Execution Matrix", level=3)
        add_table_data(doc, exp['test_matrix'])
        
    if "console_logs" in exp:
        add_styled_heading(doc, "5.2 Runtime Console Telemetry & Server Logs", level=3)
        add_code_block(doc, exp['console_logs'])
        
    add_styled_heading(doc, "5.3 Visual Verification & Execution Screenshots", level=3)
    for fig_num, caption in exp['photos']:
        add_image_placeholder(doc, fig_num, caption)
        
    # 6. Conclusion
    add_styled_heading(doc, "6. Conclusion", level=2)
    for para in exp['conclusion_text'].strip().split('\n\n'):
        if para.strip():
            add_body_paragraph(doc, para.strip())
            
    doc.add_page_break()

def generate_markdown_manual(output_path):
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("# PRACTICAL SKILLS - FULL STACK DEVELOPMENT (PS-FSD)\n")
        f.write("## COMPREHENSIVE LABORATORY MANUAL & RECORD OF EXPERIMENTS\n")
        f.write("**Academic Year:** 2026-27 | **Semester:** V  \n")
        f.write("**Enterprise Tech Stack:** Spring Boot 3.x, PostgreSQL 16, Hibernate JPA, HikariCP, SLF4J, Swagger OpenAPI, Apache Tomcat  \n")
        f.write("**Standard Format:** Aim | Objective | Theory | Procedure | Result & Discussion | Conclusion  \n\n")
        f.write("---\n\n")
        
        f.write("## Curriculum Index & Experiment Matrix\n\n")
        f.write("| Exp # | Title of Practical Experiment | Technology Focus Area | Status |\n")
        f.write("|---|---|---|---|\n")
        for exp in ALL_EXPERIMENTS:
            f.write(f"| **{exp['num']}** | [{exp['title']}](#experiment-{exp['num']}-{exp['title'].lower().replace(' ', '-').replace('/', '')}) | Spring Boot & PostgreSQL | Complete |\n")
        f.write("\n---\n\n")
        
        for exp in ALL_EXPERIMENTS:
            slug = f"experiment-{exp['num']}-{exp['title'].lower().replace(' ', '-').replace('/', '')}"
            f.write(f"## Experiment {exp['num']}: {exp['title']}\n\n")
            
            # 1. Aim
            f.write("### 1. Aim\n")
            f.write(exp['aim'].strip() + "\n\n")
            
            # 2. Objective
            f.write("### 2. Objective\n")
            for obj in exp['objective']:
                f.write(f"- {obj}\n")
            f.write("\n")
            
            # 3. Theory
            f.write("### 3. Theory\n")
            f.write(exp['theory_text'].strip() + "\n\n")
            
            if "comparison_table" in exp:
                f.write("#### Architectural Comparison & Evaluation Matrix\n\n")
                tbl = exp['comparison_table']
                f.write("| " + " | ".join(tbl[0]) + " |\n")
                f.write("| " + " | ".join(["---"] * len(tbl[0])) + " |\n")
                for row in tbl[1:]:
                    f.write("| " + " | ".join(row) + " |\n")
                f.write("\n")
                
            # 4. Procedure
            f.write("### 4. Procedure\n")
            f.write(exp['procedure_steps'].strip() + "\n\n")
            
            f.write("#### 4.1 Maven Build Configuration (`pom.xml`)\n")
            f.write("```xml\n" + exp['code_pom'].strip() + "\n```\n\n")
            
            if "code_props" in exp:
                f.write("#### 4.2 Application Configuration (`application.properties`)\n")
                f.write("```properties\n" + exp['code_props'].strip() + "\n```\n\n")
                
            if "code_main" in exp:
                f.write("#### 4.3 Java Source Code Implementation\n")
                f.write("```java\n" + exp['code_main'].strip() + "\n```\n\n")
                
            # 5. Result and Discussion
            f.write("### 5. Result and Discussion\n")
            f.write(exp['result_discussion_text'].strip() + "\n\n")
            
            if "test_matrix" in exp:
                f.write("#### 5.1 Test Execution Matrix\n\n")
                tmat = exp['test_matrix']
                f.write("| " + " | ".join(tmat[0]) + " |\n")
                f.write("| " + " | ".join(["---"] * len(tmat[0])) + " |\n")
                for row in tmat[1:]:
                    f.write("| " + " | ".join(row) + " |\n")
                f.write("\n")
                
            if "console_logs" in exp:
                f.write("#### 5.2 Runtime Console Telemetry & Server Logs\n")
                f.write("```\n" + exp['console_logs'].strip() + "\n```\n\n")
                
            f.write("#### 5.3 Visual Verification & Execution Screenshots\n\n")
            for fig_num, caption in exp['photos']:
                f.write(f"> 📸 **[SCREENSHOT PLACEHOLDER {fig_num}]**  \n")
                f.write(f"> *Figure {fig_num}: {caption}*  \n")
                f.write("> *(Paste or insert your screenshot here)*\n\n")
                
            # 6. Conclusion
            f.write("### 6. Conclusion\n")
            f.write(exp['conclusion_text'].strip() + "\n\n")
            f.write("---\n\n")

def main():
    print("Compiling Comprehensive 10-Experiment Manual (4-6 pages each)...")
    os.makedirs("submissions", exist_ok=True)
    
    doc = docx.Document()
    
    # Page setup - Standard 0.75 in margins
    for sec in doc.sections:
        sec.top_margin = Inches(0.75)
        sec.bottom_margin = Inches(0.75)
        sec.left_margin = Inches(0.75)
        sec.right_margin = Inches(0.75)
        
    # Cover Block
    cov_table = doc.add_table(rows=1, cols=1)
    cov_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = cov_table.cell(0, 0)
    set_cell_background(cell, "1B365D")
    set_cell_margins(cell, top=260, bottom=260, left=300, right=300)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING\nFULL STACK DEVELOPMENT - I (PS-FSD)\nLABORATORY MANUAL & COMPREHENSIVE EXPERIMENT RECORD\n")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("SEMESTER V | ACADEMIC YEAR 2026-27\n\nCOMPLETE 10-EXPERIMENT ENTERPRISE RECORD\nSpring Boot 3.x · PostgreSQL 16 · Hibernate JPA · REST APIs · SLF4J · Swagger OpenAPI · Apache Tomcat\n\nStandard Format: Aim | Objective | Theory | Procedure | Result & Discussion | Conclusion")
    r2.font.name = 'Arial'
    r2.font.size = Pt(9.5)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(210, 230, 250)
    
    doc.add_paragraph()
    
    add_styled_heading(doc, "Curriculum Index & Experiment Matrix", level=1)
    toc_data = [
        ["Exp #", "Title of Practical Experiment", "Technology Focus Area", "Lab Session", "Target Length"],
        ["1", "Database Creation and Configuration", "PostgreSQL DDL, HikariCP, application.properties", "Week 1", "4–6 Pages"],
        ["2", "Creating and Testing a REST Endpoint", "Spring MVC DispatcherServlet, @RestController", "Week 2", "4–6 Pages"],
        ["3", "Creating an Entity Class", "Jakarta JPA Domain Mapping, Column Constraints", "Week 3", "4–6 Pages"],
        ["4", "Creating a Repository Interface", "Spring Data JpaRepository, Dynamic Query Derivation", "Week 4", "4–6 Pages"],
        ["5", "Implementing Create and Read Operations", "POST & GET CRUD Endpoints, DTO Architecture", "Week 5", "4–6 Pages"],
        ["6", "Implementing Update and Delete Operations", "PUT & DELETE Endpoints, 404 Exception Handling", "Week 6", "4–6 Pages"],
        ["7", "Implementing Entity Relationships", "Department <-> Student (@ManyToOne / @OneToMany)", "Week 7", "4–6 Pages"],
        ["8", "Implementing Logging and Exception Handling", "SLF4J, Logback, and @RestControllerAdvice", "Week 8", "4–6 Pages"],
        ["9", "API Documentation Using Swagger/OpenAPI", "Springdoc OpenAPI 3.0 & Swagger UI Sandbox", "Week 9", "4–6 Pages"],
        ["10", "Packaging and Deployment of the Application", "WAR Build Packaging & Apache Tomcat Deployment", "Week 10", "4–6 Pages"]
    ]
    add_table_data(doc, toc_data)
    doc.add_page_break()
    
    for idx, exp in enumerate(ALL_EXPERIMENTS):
        print(f"Rendering Experiment {exp['num']}: {exp['title']}...")
        render_experiment_docx(doc, exp)
        
    docx_path = "submissions/FSD_10_Experiments_Complete_Lab_Manual.docx"
    try:
        doc.save(docx_path)
        print(f"Successfully saved Word Document: {docx_path}")
    except PermissionError:
        alt_path = "submissions/FSD_10_Experiments_Complete_Lab_Manual_Detailed.docx"
        doc.save(alt_path)
        print(f"Original file was open in Word. Successfully saved to: {alt_path}")
        root_path = "FSD_10_Experiments_Complete_Lab_Manual_Detailed.docx"
        doc.save(root_path)
        print(f"Also saved copy to root: {root_path}")
    
    md_path = "submissions/FSD_10_Experiments_Complete_Lab_Manual.md"
    generate_markdown_manual(md_path)
    print(f"Successfully saved Markdown Document: {md_path}")

if __name__ == "__main__":
    main()
