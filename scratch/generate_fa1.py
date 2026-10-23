import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F4F5F7")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    # Border
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        b = OxmlElement(f'w:{border_name}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), 'D0D7DE')
        tcBorders.append(b)
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(36, 41, 47)
    doc.add_paragraph() # space after

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    run = h.runs[0]
    if level == 1:
        run.font.name = 'Arial'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(27, 54, 93) # Navy
    elif level == 2:
        run.font.name = 'Arial'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(41, 105, 176) # Slate Blue
    elif level == 3:
        run.font.name = 'Arial'
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(51, 51, 51)
    return h

def build_fa1():
    os.makedirs('submissions', exist_ok=True)
    doc = docx.Document()
    
    # Page Margins
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)
        
    # Title Header Block
    header_table = doc.add_table(rows=1, cols=1)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = header_table.cell(0, 0)
    set_cell_background(c, "1B365D")
    set_cell_margins(c, top=200, bottom=200, left=250, right=250)
    
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("FULL STACK DEVELOPMENT - I (PS-FSD)\nFORMATIVE ASSIGNMENT - 1 (FA-1)\nSEM V | A.Y. 2026-27")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    
    p2 = c.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("PROJECT: GrowthOS — The AI Growth Operating System for E-Commerce Brands\nBloom's Taxonomy Levels: L4 (Analyze) & L6 (Create)")
    r2.font.name = 'Arial'
    r2.font.size = Pt(10)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(200, 220, 245)
    
    doc.add_paragraph()
    
    # Project Overview Box
    add_styled_heading(doc, "Executive Summary & Project Context", level=1)
    p = doc.add_paragraph(
        "GrowthOS is a production-grade full-stack SaaS platform designed for direct-to-consumer (D2C) e-commerce brands. "
        "The system aggregates core business metrics (Revenue, COGS, Blended CAC, ROAS, RTO, Cart Abandonment), executes a parallel 7-agent "
        "LLM diagnostic pipeline, runs Monte Carlo business strategy simulations, and generates executive CEO reports with actionable SOPs. "
        "This submission presents the comprehensive database design, Entity-Relationship (ER) model, entity attribute breakdown, "
        "relationship cardinality analysis, and relational integrity validation for GrowthOS."
    )
    p.paragraph_format.line_spacing = 1.15
    
    # Task a
    add_styled_heading(doc, "Task (a): Entity Identification & Attributes Specification", level=1)
    p = doc.add_paragraph(
        "The GrowthOS architecture is modeled with normalized relational tables backed by PostgreSQL and managed through "
        "Spring Data JPA / Prisma ORM. Below is the comprehensive breakdown of primary entities, identifying Primary Keys (PK), "
        "Foreign Keys (FK), data types, constraints, and descriptions."
    )
    p.paragraph_format.line_spacing = 1.15
    
    # Entity 1: User
    add_styled_heading(doc, "1. User Entity (Authentication & Tenant Owner)", level=2)
    user_attrs = [
        ["Attribute", "Data Type", "Constraint", "Description"],
        ["id", "VARCHAR(36) / String", "PK (Primary Key, CUID/UUID)", "Unique identifier for the user account"],
        ["email", "VARCHAR(255)", "UNIQUE, NOT NULL", "Login email address and primary identity identifier"],
        ["name", "VARCHAR(255)", "NULLABLE", "Full name of the business founder / user"],
        ["password", "VARCHAR(255)", "NULLABLE", "Bcrypt-hashed credential password for local auth"],
        ["emailVerified", "TIMESTAMP", "NULLABLE", "Verification timestamp for email confirmation"],
        ["image", "TEXT / VARCHAR(500)", "NULLABLE", "Avatar / profile photo URL from OAuth provider"],
        ["createdAt", "TIMESTAMP", "DEFAULT NOW(), NOT NULL", "Account creation audit timestamp"],
        ["updatedAt", "TIMESTAMP", "DEFAULT NOW(), NOT NULL", "Last profile update audit timestamp"]
    ]
    t = doc.add_table(rows=len(user_attrs), cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(user_attrs):
        for c_idx, val in enumerate(row):
            cell = t.cell(r_idx, c_idx)
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Arial'
            if r_idx == 0:
                set_cell_background(cell, "2969B0")
                run.font.size = Pt(9.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "F8F9FA" if r_idx % 2 == 1 else "FFFFFF")
                run.font.size = Pt(8.5)
                if c_idx == 0 or "PK" in val or "FK" in val:
                    run.font.bold = True
    doc.add_paragraph()
    
    # Entity 2: Business
    add_styled_heading(doc, "2. Business Entity (E-Commerce Store & Operational Metrics)", level=2)
    biz_attrs = [
        ["Attribute", "Data Type", "Constraint", "Description"],
        ["id", "VARCHAR(36)", "PK", "Unique store identifier (CUID)"],
        ["userId", "VARCHAR(36)", "FK -> User(id), NOT NULL", "Owner reference; ON DELETE CASCADE"],
        ["name", "VARCHAR(255)", "NOT NULL", "Store / brand name (e.g., Ethnic Elegance)"],
        ["businessType", "ENUM / VARCHAR(50)", "NOT NULL, DEFAULT ECOMMERCE", "Category (ECOMMERCE, RETAIL, D2C, etc.)"],
        ["industry", "VARCHAR(100)", "NULLABLE", "Sub-vertical (Apparel, Beauty, Electronics)"],
        ["website", "VARCHAR(255)", "NULLABLE", "Storefront URL"],
        ["monthlyRevenue", "DECIMAL(14,2)", "NULLABLE", "Gross monthly sales in INR"],
        ["monthlyExpenses", "DECIMAL(14,2)", "NULLABLE", "Total monthly operating costs (COGS + Opex)"],
        ["customersPerMonth", "INTEGER", "NULLABLE", "Monthly transacting unique buyer count"],
        ["repeatCustomerPercentage", "INTEGER", "CHECK (0-100)", "Percentage of orders from repeat buyers"],
        ["marketingChannels", "TEXT[] / VARCHAR[]", "DEFAULT '{}'", "Active acquisition channels (Meta, Google, etc.)"],
        ["pricingLevel", "ENUM(LOW, MEDIUM, HIGH)", "NOT NULL", "Market positioning bracket"],
        ["healthScore", "INTEGER", "CHECK (0-100)", "Composite AI health benchmark score (0-100)"],
        ["averageOrderValue", "DECIMAL(10,2)", "NULLABLE", "Blended AOV in INR"],
        ["customerAcquisitionCost", "DECIMAL(10,2)", "NULLABLE", "Blended CAC in INR"],
        ["returnOnAdSpend", "DECIMAL(6,2)", "NULLABLE", "ROAS multiplier (e.g., 2.80x)"],
        ["cartAbandonmentRate", "INTEGER", "CHECK (0-100)", "Checkout drop-off rate %"],
        ["inventoryTurnover", "DECIMAL(6,2)", "NULLABLE", "Annual inventory turns ratio"],
        ["fulfillmentMethod", "ENUM", "NULLABLE", "SELF_FULFILLED, THREE_PL, DROPSHIPPING"],
        ["codRate", "INTEGER", "CHECK (0-100)", "Cash on Delivery order share %"],
        ["rtoRate", "INTEGER", "CHECK (0-100)", "Return to Origin failure rate %"],
        ["prepaidPercentage", "INTEGER", "CHECK (0-100)", "Prepaid transaction share %"],
        ["createdAt", "TIMESTAMP", "NOT NULL, DEFAULT NOW()", "Store profile creation timestamp"]
    ]
    t = doc.add_table(rows=len(biz_attrs), cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(biz_attrs):
        for c_idx, val in enumerate(row):
            cell = t.cell(r_idx, c_idx)
            set_cell_margins(cell, top=50, bottom=50, left=90, right=90)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Arial'
            if r_idx == 0:
                set_cell_background(cell, "2969B0")
                run.font.size = Pt(9.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "F8F9FA" if r_idx % 2 == 1 else "FFFFFF")
                run.font.size = Pt(8.0)
                if c_idx == 0 or "PK" in val or "FK" in val:
                    run.font.bold = True
    doc.add_paragraph()

    # Entity 3: BusinessSnapshot
    add_styled_heading(doc, "3. BusinessSnapshot Entity (Time-Series Metric History)", level=2)
    snap_attrs = [
        ["Attribute", "Data Type", "Constraint", "Description"],
        ["id", "VARCHAR(36)", "PK", "Snapshot record ID"],
        ["businessId", "VARCHAR(36)", "FK -> Business(id), NOT NULL", "Referenced store; ON DELETE CASCADE"],
        ["revenue", "DECIMAL(14,2)", "NULLABLE", "Revenue captured at point in time"],
        ["expenses", "DECIMAL(14,2)", "NULLABLE", "Expenses captured at point in time"],
        ["customers", "INTEGER", "NULLABLE", "Customer count at snapshot"],
        ["healthScore", "INTEGER", "NULLABLE", "Historical composite health score (0-100)"],
        ["averageOrderValue", "DECIMAL(10,2)", "NULLABLE", "Historical AOV"],
        ["customerAcquisitionCost", "DECIMAL(10,2)", "NULLABLE", "Historical CAC"],
        ["returnOnAdSpend", "DECIMAL(6,2)", "NULLABLE", "Historical ROAS"],
        ["codRate", "INTEGER", "NULLABLE", "Historical COD %"],
        ["rtoRate", "INTEGER", "NULLABLE", "Historical RTO %"],
        ["createdAt", "TIMESTAMP", "NOT NULL, DEFAULT NOW()", "Snapshot capture timestamp (supports time-series analysis)"]
    ]
    t = doc.add_table(rows=len(snap_attrs), cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(snap_attrs):
        for c_idx, val in enumerate(row):
            cell = t.cell(r_idx, c_idx)
            set_cell_margins(cell, top=50, bottom=50, left=90, right=90)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Arial'
            if r_idx == 0:
                set_cell_background(cell, "2969B0")
                run.font.size = Pt(9.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "F8F9FA" if r_idx % 2 == 1 else "FFFFFF")
                run.font.size = Pt(8.0)
                if c_idx == 0 or "PK" in val or "FK" in val:
                    run.font.bold = True
    doc.add_paragraph()

    # Entity 4: Analysis
    add_styled_heading(doc, "4. Analysis Entity (7-Agent AI Diagnostic Execution Run)", level=2)
    analysis_attrs = [
        ["Attribute", "Data Type", "Constraint", "Description"],
        ["id", "VARCHAR(36)", "PK", "Diagnosis run identifier"],
        ["businessId", "VARCHAR(36)", "FK -> Business(id), NOT NULL", "Store being analyzed; ON DELETE CASCADE"],
        ["createdById", "VARCHAR(36)", "FK -> User(id), NULLABLE", "User initiating run; ON DELETE SET NULL"],
        ["status", "ENUM", "PENDING, IN_PROGRESS, COMPLETED, FAILED", "Execution state of the 7-agent pipeline"],
        ["healthScore", "INTEGER", "CHECK (0-100)", "Evaluated health rating from 7 agents"],
        ["summary", "JSONB", "NULLABLE", "Synthesized executive findings across agents"],
        ["agents", "JSONB", "NULLABLE", "Structured output per agent (Finance, Marketing, Ops, Tech, etc.)"],
        ["inputSnapshot", "JSONB", "NULLABLE", "Frozen inputs at execution time for auditability"],
        ["createdAt", "TIMESTAMP", "NOT NULL, DEFAULT NOW()", "Execution start timestamp"],
        ["updatedAt", "TIMESTAMP", "NOT NULL, DEFAULT NOW()", "Last status transition timestamp"]
    ]
    t = doc.add_table(rows=len(analysis_attrs), cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(analysis_attrs):
        for c_idx, val in enumerate(row):
            cell = t.cell(r_idx, c_idx)
            set_cell_margins(cell, top=50, bottom=50, left=90, right=90)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Arial'
            if r_idx == 0:
                set_cell_background(cell, "2969B0")
                run.font.size = Pt(9.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "F8F9FA" if r_idx % 2 == 1 else "FFFFFF")
                run.font.size = Pt(8.0)
                if c_idx == 0 or "PK" in val or "FK" in val:
                    run.font.bold = True
    doc.add_paragraph()

    # Entity 5: Simulation
    add_styled_heading(doc, "5. Simulation Entity (What-If Strategy Forecasting)", level=2)
    sim_attrs = [
        ["Attribute", "Data Type", "Constraint", "Description"],
        ["id", "VARCHAR(36)", "PK", "Simulation experiment ID"],
        ["businessId", "VARCHAR(36)", "FK -> Business(id), NOT NULL", "Store reference; ON DELETE CASCADE"],
        ["createdById", "VARCHAR(36)", "FK -> User(id), NULLABLE", "User reference; ON DELETE SET NULL"],
        ["status", "ENUM", "PENDING, COMPLETED, FAILED", "Simulation job status"],
        ["strategy", "TEXT", "NOT NULL", "Proposed hypothesis (e.g., 'Cut ad spend 20%, shift to WhatsApp')"],
        ["shortTermOutcome", "JSONB", "NULLABLE", "30-90 day forecasted metrics (Revenue, ROAS, RTO)"],
        ["longTermOutcome", "JSONB", "NULLABLE", "6-12 month strategic trajectory projection"],
        ["risks", "JSONB", "NULLABLE", "Identified failure modes and mitigating controls"],
        ["opportunities", "JSONB", "NULLABLE", "Upside potential and leverage points"],
        ["confidenceScore", "INTEGER", "CHECK (0-100)", "Model confidence estimate based on input density"],
        ["createdAt", "TIMESTAMP", "NOT NULL, DEFAULT NOW()", "Execution timestamp"]
    ]
    t = doc.add_table(rows=len(sim_attrs), cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(sim_attrs):
        for c_idx, val in enumerate(row):
            cell = t.cell(r_idx, c_idx)
            set_cell_margins(cell, top=50, bottom=50, left=90, right=90)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Arial'
            if r_idx == 0:
                set_cell_background(cell, "2969B0")
                run.font.size = Pt(9.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "F8F9FA" if r_idx % 2 == 1 else "FFFFFF")
                run.font.size = Pt(8.0)
                if c_idx == 0 or "PK" in val or "FK" in val:
                    run.font.bold = True
    doc.add_paragraph()

    # Entity 6: Report
    add_styled_heading(doc, "6. Report Entity (Executive Briefing & Exportable Artifacts)", level=2)
    rep_attrs = [
        ["Attribute", "Data Type", "Constraint", "Description"],
        ["id", "VARCHAR(36)", "PK", "Report document identifier"],
        ["businessId", "VARCHAR(36)", "FK -> Business(id), NOT NULL", "Store target; ON DELETE CASCADE"],
        ["createdById", "VARCHAR(36)", "FK -> User(id), NULLABLE", "Author reference; ON DELETE SET NULL"],
        ["analysisId", "VARCHAR(36)", "FK -> Analysis(id), NULLABLE", "Linked diagnosis run; ON DELETE SET NULL"],
        ["simulationId", "VARCHAR(36)", "FK -> Simulation(id), NULLABLE", "Linked simulation run; ON DELETE SET NULL"],
        ["type", "ENUM", "DIAGNOSIS, SIMULATION, EXECUTIVE_SUMMARY", "Report category"],
        ["title", "VARCHAR(255)", "NOT NULL", "Report title header"],
        ["executiveSummary", "TEXT", "NULLABLE", "High-level briefing for founders / investors"],
        ["chiefRisk", "TEXT", "NULLABLE", "Single most pressing threat to store profitability"],
        ["topOpportunity", "TEXT", "NULLABLE", "Highest ROI lever identified"],
        ["immediateAction", "TEXT", "NULLABLE", "First 7-day tactical priority"],
        ["actionPlan", "JSONB", "NULLABLE", "Structured step-by-step roadmap with assigned tools & deadlines"],
        ["createdAt", "TIMESTAMP", "NOT NULL, DEFAULT NOW()", "Generation timestamp"]
    ]
    t = doc.add_table(rows=len(rep_attrs), cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(rep_attrs):
        for c_idx, val in enumerate(row):
            cell = t.cell(r_idx, c_idx)
            set_cell_margins(cell, top=50, bottom=50, left=90, right=90)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Arial'
            if r_idx == 0:
                set_cell_background(cell, "2969B0")
                run.font.size = Pt(9.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "F8F9FA" if r_idx % 2 == 1 else "FFFFFF")
                run.font.size = Pt(8.0)
                if c_idx == 0 or "PK" in val or "FK" in val:
                    run.font.bold = True
    doc.add_paragraph()

    # Auxiliary Entities
    add_styled_heading(doc, "7. Auxiliary & AI System Entities", level=2)
    aux_p = doc.add_paragraph(
        "- AgentFeedback: id (PK), analysisId (FK), agentName (VARCHAR), rating (INT 1-5), comment (TEXT), userId (FK), createdAt (TIMESTAMP).\n"
        "- KnowledgeChunk (RAG): id (PK), content (TEXT), source (VARCHAR), category (VARCHAR), industry (VARCHAR, NULLABLE), createdAt (TIMESTAMP).\n"
        "- Account / Session: Standard NextAuth security adapter entities managing OAuth tokens and browser sessions with cascade deletion linked to User."
    )
    aux_p.paragraph_format.line_spacing = 1.15
    doc.add_paragraph()

    # Task b
    add_styled_heading(doc, "Task (b): Relational Analysis & Cardinality Specification", level=1)
    p = doc.add_paragraph(
        "The relational schema of GrowthOS adheres to rigorous normalization (3NF) to eliminate data redundancy while "
        "leveraging PostgreSQL JSONB data types for dynamic, agentic AI outputs. Below is the analytical breakdown of entity associations:"
    )
    p.paragraph_format.line_spacing = 1.15
    
    rel_rows = [
        ["Relationship", "Parent Entity", "Child Entity", "Cardinality", "Foreign Key", "Cascade Action", "Business Rationale"],
        ["User owns Businesses", "User", "Business", "1 : N (One-to-Many)", "Business.userId", "ON DELETE CASCADE", "One founder can manage multiple e-commerce brand storefronts."],
        ["Business has Snapshots", "Business", "BusinessSnapshot", "1 : N (One-to-Many)", "BusinessSnapshot.businessId", "ON DELETE CASCADE", "Periodic metric captures enable longitudinal tracking and historical charting."],
        ["Business has Diagnoses", "Business", "Analysis", "1 : N (One-to-Many)", "Analysis.businessId", "ON DELETE CASCADE", "A store runs multiple 7-agent AI diagnoses over time."],
        ["User initiates Analysis", "User", "Analysis", "1 : N (One-to-Many)", "Analysis.createdById", "ON DELETE SET NULL", "Audit log tracks who requested the run; retains report if user deleted."],
        ["Business runs Simulations", "Business", "Simulation", "1 : N (One-to-Many)", "Simulation.businessId", "ON DELETE CASCADE", "A brand simulates various strategic scenarios (e.g., cutting ad budget)."],
        ["User initiates Simulation", "User", "Simulation", "1 : N (One-to-Many)", "Simulation.createdById", "ON DELETE SET NULL", "Audit trail of scenario execution."],
        ["Business has Reports", "Business", "Report", "1 : N (One-to-Many)", "Report.businessId", "ON DELETE CASCADE", "Consolidated CEO reports generated per store."],
        ["Analysis generates Report", "Analysis", "Report", "1 : N (One-to-Many)", "Report.analysisId", "ON DELETE SET NULL", "A diagnosis run can be exported into one or more executive PDF reports."],
        ["Simulation generates Report", "Simulation", "Report", "1 : N (One-to-Many)", "Report.simulationId", "ON DELETE SET NULL", "Simulation projections can be compiled into strategic reports."]
    ]
    t = doc.add_table(rows=len(rel_rows), cols=7)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(rel_rows):
        for c_idx, val in enumerate(row):
            cell = t.cell(r_idx, c_idx)
            set_cell_margins(cell, top=50, bottom=50, left=70, right=70)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Arial'
            if r_idx == 0:
                set_cell_background(cell, "2969B0")
                run.font.size = Pt(8.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "F8F9FA" if r_idx % 2 == 1 else "FFFFFF")
                run.font.size = Pt(7.5)
                if c_idx in [0, 3]:
                    run.font.bold = True
    doc.add_paragraph()

    # Task c: Complete ER Diagram
    add_styled_heading(doc, "Task (c): Complete Entity-Relationship (ER) Diagram", level=1)
    p = doc.add_paragraph(
        "Below is the complete architectural Entity-Relationship Diagram representing the GrowthOS database schema, "
        "rendered via Crow's Foot notation and Mermaid formal grammar. It illustrates all primary keys (PK), foreign keys (FK), "
        "relational links, and multiplicity cardinalities."
    )
    p.paragraph_format.line_spacing = 1.15

    # Visual Crow's foot diagram text box
    er_diagram_ascii = """
+-------------------------------------------------------------------------------------------------------+
|                                      GROWTHOS ER DIAGRAM OVERVIEW                                      |
+-------------------------------------------------------------------------------------------------------+

        +-----------------------+                    1 : N
        |         USER          | -------------------------------------------+
        |-----------------------|                                            |
        | PK  id                |                                            |
        |     email (UQ)        |                                            v
        |     name              | 1 : N                             +-----------------------+
        |     password          | -------------------------------+  |       BUSINESS        |
        |     createdAt         |                                |  |-----------------------|
        +-----------------------+                                |  | PK  id                |
                   |                                             |  | FK  userId            |
                   | 1 : N                                       |  |     name              |
                   v                                             |  |     businessType      |
        +-----------------------+                                |  |     monthlyRevenue    |
        |   ACCOUNT / SESSION   |                                |  |     monthlyExpenses   |
        |-----------------------|                                |  |     codRate / rtoRate |
        | PK  id                |                                |  |     healthScore       |
        | FK  userId            |                                |  +-----------------------+
        +-----------------------+                                |         |     |     |
                                                                 |         |     |     | 1 : N
              +--------------------------------------------------+         |     |     +------------+
              |                                1 : N                       |     | 1 : N            |
              |       +----------------------------------------------------+     v                  v
              |       |                                               +------------------+ +------------------+
              v       v                                               | BUSINESS_SNAPSHOT| |    SIMULATION    |
        +-----------------------+                                     |------------------| |------------------|
        |       ANALYSIS        |                                     | PK  id           | | PK  id           |
        |-----------------------|                                     | FK  businessId   | | FK  businessId   |
        | PK  id                |                                     |     revenue      | | FK  createdById  |
        | FK  businessId        |                                     |     expenses     | |     strategy     |
        | FK  createdById       |                                     |     healthScore  | |     shortOutcome |
        |     status            |                                     |     createdAt    | |     risks/opps   |
        |     healthScore       |                                     +------------------+ +------------------+
        |     agents (JSONB)    |                                                                   |
        |     summary (JSONB)   |                                                                   | 1 : N
        +-----------------------+                                                                   |
                   |                                                                                |
                   | 1 : N                                                                          |
                   v                                                                                v
        +-----------------------------------------------------------------------------------------------+
        |                                            REPORT                                             |
        |-----------------------------------------------------------------------------------------------|
        | PK  id                                                                                        |
        | FK  businessId  (-> Business.id, CASCADE)                                                      |
        | FK  createdById (-> User.id, SET NULL)                                                        |
        | FK  analysisId  (-> Analysis.id, SET NULL)                                                    |
        | FK  simulationId(-> Simulation.id, SET NULL)                                                  |
        |     type, title, executiveSummary, chiefRisk, topOpportunity, actionPlan (JSONB), createdAt   |
        +-----------------------------------------------------------------------------------------------+
    """
    add_code_block(doc, er_diagram_ascii)

    add_styled_heading(doc, "Mermaid ER Grammar Specification", level=2)
    mermaid_spec = """
erDiagram
    USER ||--o{ BUSINESS : "owns (1:N)"
    USER ||--o{ ANALYSIS : "triggers (1:N)"
    USER ||--o{ SIMULATION : "runs (1:N)"
    USER ||--o{ REPORT : "generates (1:N)"
    USER ||--o{ ACCOUNT : "authenticates (1:N)"
    USER ||--o{ SESSION : "maintains (1:N)"

    BUSINESS ||--o{ BUSINESS_SNAPSHOT : "captures (1:N)"
    BUSINESS ||--o{ ANALYSIS : "diagnosed_by (1:N)"
    BUSINESS ||--o{ SIMULATION : "simulated_by (1:N)"
    BUSINESS ||--o{ REPORT : "reported_in (1:N)"

    ANALYSIS ||--o{ REPORT : "summarized_in (1:N)"
    SIMULATION ||--o{ REPORT : "projected_in (1:N)"

    USER {
        string id PK
        string email UK
        string name
        string password
        timestamp createdAt
    }

    BUSINESS {
        string id PK
        string userId FK
        string name
        string businessType
        decimal monthlyRevenue
        decimal monthlyExpenses
        int healthScore
        int codRate
        int rtoRate
    }

    BUSINESS_SNAPSHOT {
        string id PK
        string businessId FK
        decimal revenue
        decimal expenses
        int healthScore
        timestamp createdAt
    }

    ANALYSIS {
        string id PK
        string businessId FK
        string createdById FK
        string status
        int healthScore
        jsonb agents
        jsonb summary
    }

    SIMULATION {
        string id PK
        string businessId FK
        string createdById FK
        string status
        string strategy
        jsonb shortTermOutcome
        jsonb risks
    }

    REPORT {
        string id PK
        string businessId FK
        string createdById FK
        string analysisId FK
        string simulationId FK
        string type
        string title
        text executiveSummary
        jsonb actionPlan
    }
    """
    add_code_block(doc, mermaid_spec)

    # Task d
    add_styled_heading(doc, "Task (d): Consistency Verification & Architectural Justification", level=1)
    p = doc.add_paragraph(
        "A rigorous consistency check was performed comparing the conceptual ER diagram, the physical PostgreSQL database schema, "
        "and the Object-Relational Mapping (ORM) entity models in Spring Data JPA / Prisma. The table below documents the "
        "architectural decisions and explains how discrepancies between pure relational theory and real-world high-throughput applications were resolved:"
    )
    p.paragraph_format.line_spacing = 1.15

    disc_rows = [
        ["Architectural Dimension", "Pure Relational (Academic)", "GrowthOS PostgreSQL / ORM Design", "Engineering Justification"],
        ["Primary Key Strategy", "Auto-increment BIGSERIAL integer", "CUID / UUID (String PK)", "Mitigates sequential ID enumeration attacks; permits client-side ID pre-generation; multi-region distributed safe."],
        ["Multi-Agent Diagnostics Storage", "Dozens of normalized 1:N tables (Agent, Problem, Action)", "JSONB Columns (agents, summary, actionPlan)", "High read throughput for 7 parallel agents; sub-4s response; allows flexible agent schemas without running DDL migrations on every prompt update."],
        ["Time-Series Auditing", "Mutable columns on parent store table", "Immutable BusinessSnapshot table + composite index", "Preserves exact historical state for interactive trend lines and regression analysis."],
        ["Financial Computations", "DOUBLE / FLOAT floating-point types", "DECIMAL(14,2) with strict precision", "Eliminates IEEE-754 binary floating-point rounding errors in revenue, COGS, and profit calculations."],
        ["Referential Integrity & Cascades", "Universal RESTRICT constraints", "Selective CASCADE vs SET NULL", "Deleting a store purges its snapshots and analyses (CASCADE); deleting a team member preserves generated business reports (SET NULL)."]
    ]
    t = doc.add_table(rows=len(disc_rows), cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(disc_rows):
        for c_idx, val in enumerate(row):
            cell = t.cell(r_idx, c_idx)
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Arial'
            if r_idx == 0:
                set_cell_background(cell, "2969B0")
                run.font.size = Pt(8.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "F8F9FA" if r_idx % 2 == 1 else "FFFFFF")
                run.font.size = Pt(7.5)
                if c_idx == 0:
                    run.font.bold = True
    doc.add_paragraph()

    # Spring Boot JPA Entity Classes Mapping
    add_styled_heading(doc, "Appendix: Spring Boot JPA Entity Mapping Reference", level=1)
    p = doc.add_paragraph(
        "To demonstrate full compliance with Spring Data JPA annotations as required in the course curriculum, "
        "the standard JPA Entity implementations for the core GrowthOS entities (`User.java` and `Business.java`) are provided below:"
    )
    p.paragraph_format.line_spacing = 1.15

    jpa_sample = """
// User.java - Spring Boot JPA Entity
package com.growthos.entity;

import jakarta.persistence.*;
import lombok.*;
import java.time.LocalDateTime;
import java.util.List;

@Entity
@Table(name = "users")
@Getter @Setter @NoArgsConstructor @AllArgsConstructor @Builder
public class User {
    @Id
    @Column(length = 36)
    private String id;

    @Column(nullable = false, unique = true, length = 255)
    private String email;

    @Column(length = 255)
    private String name;

    @Column(length = 255)
    private String password;

    @OneToMany(mappedBy = "user", cascade = CascadeType.ALL, orphanRemoval = true)
    private List<Business> businesses;

    @Column(nullable = false, updatable = false)
    private LocalDateTime createdAt = LocalDateTime.now();
}

// Business.java - Spring Boot JPA Entity
package com.growthos.entity;

import jakarta.persistence.*;
import lombok.*;
import java.math.BigDecimal;
import java.time.LocalDateTime;
import java.util.List;

@Entity
@Table(name = "businesses", indexes = {
    @Index(name = "idx_business_user", columnList = "user_id"),
    @Index(name = "idx_business_type", columnList = "businessType")
})
@Getter @Setter @NoArgsConstructor @AllArgsConstructor @Builder
public class Business {
    @Id
    @Column(length = 36)
    private String id;

    @ManyToOne(fetch = FetchType.LAZY)
    @JoinColumn(name = "user_id", nullable = false)
    private User user;

    @Column(nullable = false, length = 255)
    private String name;

    @Enumerated(EnumType.STRING)
    @Column(nullable = false)
    private BusinessType businessType;

    @Column(precision = 14, scale = 2)
    private BigDecimal monthlyRevenue;

    @Column(precision = 14, scale = 2)
    private BigDecimal monthlyExpenses;

    private Integer healthScore;
    private Integer codRate;
    private Integer rtoRate;

    @OneToMany(mappedBy = "business", cascade = CascadeType.ALL)
    private List<BusinessSnapshot> snapshots;

    @OneToMany(mappedBy = "business", cascade = CascadeType.ALL)
    private List<Analysis> analyses;

    @Column(nullable = false, updatable = false)
    private LocalDateTime createdAt = LocalDateTime.now();
}
"""
    add_code_block(doc, jpa_sample)

    output_path = "submissions/FA1_GrowthOS_ER_Diagram_Submission.docx"
    doc.save(output_path)
    print(f"Successfully generated: {output_path}")

build_fa1()
