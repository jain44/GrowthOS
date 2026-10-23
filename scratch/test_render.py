import sys
import os

current_dir = os.path.dirname(os.path.abspath(__file__))
project_root = os.path.abspath(os.path.join(current_dir, ".."))
os.chdir(project_root)

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from experiments_common import (
    set_cell_background, set_cell_margins, add_image_placeholder,
    add_code_block, add_table_data, add_styled_heading, add_body_paragraph
)

from exp_data_part1 import EXP_1

def render_experiment(doc, exp_data):
    add_styled_heading(doc, f"Experiment {exp_data['num']}: {exp_data['title']}", level=1)
    
    # Section 1: Aim
    add_styled_heading(doc, "1. Aim", level=2)
    add_body_paragraph(doc, exp_data['aim'])
    
    # Section 2: Objective
    add_styled_heading(doc, "2. Objective", level=2)
    for obj in exp_data['objective']:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(obj)
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(40, 40, 40)
    doc.add_paragraph()
    
    # Section 3: Theory
    add_styled_heading(doc, "3. Theory", level=2)
    for para in exp_data['theory_text'].strip().split('\n\n'):
        if para.strip():
            add_body_paragraph(doc, para.strip())
            
    if "comparison_table" in exp_data:
        add_styled_heading(doc, "Architectural Comparison & Evaluation Matrix", level=3)
        add_table_data(doc, exp_data['comparison_table'])
        
    # Section 4: Procedure
    add_styled_heading(doc, "4. Procedure", level=2)
    for para in exp_data['procedure_steps'].strip().split('\n\n'):
        if para.strip():
            add_body_paragraph(doc, para.strip())
            
    add_styled_heading(doc, "4.1 Maven Build Configuration (pom.xml)", level=3)
    add_code_block(doc, exp_data['code_pom'])
    
    if "code_props" in exp_data:
        add_styled_heading(doc, "4.2 Application Configuration (src/main/resources/application.properties)", level=3)
        add_code_block(doc, exp_data['code_props'])
        
    if "code_main" in exp_data:
        add_styled_heading(doc, "4.3 Java Source Code Implementation", level=3)
        add_code_block(doc, exp_data['code_main'])
        
    # Section 5: Result and Discussion
    add_styled_heading(doc, "5. Result and Discussion", level=2)
    for para in exp_data['result_discussion_text'].strip().split('\n\n'):
        if para.strip():
            add_body_paragraph(doc, para.strip())
            
    if "test_matrix" in exp_data:
        add_styled_heading(doc, "5.1 Test Execution Matrix", level=3)
        add_table_data(doc, exp_data['test_matrix'])
        
    if "console_logs" in exp_data:
        add_styled_heading(doc, "5.2 Runtime Console Telemetry & Server Logs", level=3)
        add_code_block(doc, exp_data['console_logs'])
        
    add_styled_heading(doc, "5.3 Visual Verification & Execution Screenshots", level=3)
    for fig_num, caption in exp_data['photos']:
        add_image_placeholder(doc, fig_num, caption)
        
    # Section 6: Conclusion
    add_styled_heading(doc, "6. Conclusion", level=2)
    for para in exp_data['conclusion_text'].strip().split('\n\n'):
        if para.strip():
            add_body_paragraph(doc, para.strip())
            
    doc.add_page_break()

print("Renderer ready.")
