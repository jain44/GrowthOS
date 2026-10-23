import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from experiments_builder import (
    set_cell_background, set_cell_margins, add_image_placeholder,
    add_code_block, add_table_data, add_styled_heading, add_body_paragraph
)

def create_massive_manual():
    doc = docx.Document()
    
    # Page setup - Standard Letter / A4 with 0.75 in margins
    for sec in doc.sections:
        sec.top_margin = Inches(0.75)
        sec.bottom_margin = Inches(0.75)
        sec.left_margin = Inches(0.75)
        sec.right_margin = Inches(0.75)
        
    # =========================================================================
    # COVER PAGE
    # =========================================================================
    cov_table = doc.add_table(rows=1, cols=1)
    cov_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = cov_table.cell(0, 0)
    set_cell_background(cell, "1B365D")
    set_cell_margins(cell, top=260, bottom=260, left=300, right=300)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING\nFULL STACK DEVELOPMENT - I (PS-FSD)\nLABORATORY MANUAL & COMPREHENSIVE EXPERIMENT RECORD\n")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("SEMESTER V | ACADEMIC YEAR 2026-27\n\nCOMPLETE 10-EXPERIMENT ENTERPRISE MANUAL\nSpring Boot 3.x · PostgreSQL 16 · Hibernate JPA · REST APIs · SLF4J · Swagger OpenAPI · Apache Tomcat\n\nStandard Format: Aim | Objective | Theory | Procedure | Result & Discussion | Conclusion")
    r2.font.name = 'Arial'
    r2.font.size = Pt(10)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(210, 230, 250)
    
    doc.add_paragraph()
    
    # Table of Contents
    add_styled_heading(doc, "Curriculum Index & Experiment Matrix", level=1)
    toc_data = [
        ["Exp #", "Title of Practical Experiment", "Technology Focus Area", "Lab Session", "Target Pages"],
        ["1", "Database Creation and Configuration", "PostgreSQL DDL, HikariCP, application.properties", "Week 1", "pp. 2–6"],
        ["2", "Creating and Testing a REST Endpoint", "Spring MVC DispatcherServlet, @RestController", "Week 2", "pp. 7–12"],
        ["3", "Creating an Entity Class", "Jakarta JPA Domain Mapping, Column Constraints", "Week 3", "pp. 13–18"],
        ["4", "Creating a Repository Interface", "Spring Data JpaRepository, Dynamic Query Derivation", "Week 4", "pp. 19–24"],
        ["5", "Implementing Create and Read Operations", "POST & GET CRUD Endpoints, DTO Architecture", "Week 5", "pp. 25–30"],
        ["6", "Implementing Update and Delete Operations", "PUT & DELETE Endpoints, 404 Exception Handling", "Week 6", "pp. 31–36"],
        ["7", "Implementing Entity Relationships", "Department <-> Student (@ManyToOne / @OneToMany)", "Week 7", "pp. 37–43"],
        ["8", "Implementing Logging and Exception Handling", "SLF4J, Logback, and @RestControllerAdvice", "Week 8", "pp. 44–50"],
        ["9", "API Documentation Using Swagger/OpenAPI", "Springdoc OpenAPI 3.0 & Swagger UI Sandbox", "Week 9", "pp. 51–56"],
        ["10", "Packaging and Deployment of the Application", "WAR Build Packaging & Apache Tomcat Deployment", "Week 10", "pp. 57–62"]
    ]
    add_table_data(doc, toc_data)
    doc.add_page_break()

    print("Cover and index generated.")

create_massive_manual()
