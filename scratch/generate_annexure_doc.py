import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

current_dir = os.path.dirname(os.path.abspath(__file__))
project_root = os.path.abspath(os.path.join(current_dir, ".."))
os.chdir(project_root)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_heading_16(doc, text):
    """Headings: 16 Bold, Times New Roman"""
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(27, 54, 93) # Deep Navy
    return h

def add_content_heading_14(doc, text):
    """Content Headings: 14 Bold, Times New Roman"""
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(35, 75, 125) # Slate Navy
    return h

def add_content_12(doc, text, bold_prefix=None):
    """Content: 12, Times New Roman"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(20, 20, 20)
    
    r_body = p.add_run(text)
    r_body.font.name = 'Times New Roman'
    r_body.font.size = Pt(12)
    r_body.font.color.rgb = RGBColor(40, 40, 40)
    return p

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        b = OxmlElement(f'w:{border_name}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), 'D0D7DE')
        tcBorders.append(b)
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text.strip())
    run.font.name = 'Consolas'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 30, 30)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(2)
    p_spacer.paragraph_format.space_after = Pt(4)

def read_file(rel_path):
    path = os.path.join(project_root, rel_path)
    if os.path.exists(path):
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    return f"// File not found: {rel_path}"

def build_annexure_doc():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Header / Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_t = title_p.add_run("GROWTHOS — FSD-1 PROJECT REPORT")
    run_t.font.name = 'Times New Roman'
    run_t.font.size = Pt(18)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(27, 54, 93)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(2)
    sub_p.paragraph_format.space_after = Pt(14)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_s = sub_p.add_run("SECTION 7: ANNEXURE (COMPLETE SOURCE CODE & ARCHITECTURE IMPLEMENTATION)")
    run_s.font.name = 'Times New Roman'
    run_s.font.size = Pt(14)
    run_s.font.bold = True
    run_s.font.color.rgb = RGBColor(70, 70, 70)

    # Metadata Table
    meta_tbl = doc.add_table(rows=4, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Course & Project Track:", "Full Stack Development - 1 (FSD-1) Academic Capstone"),
        ("Project System Name:", "GrowthOS — Autonomous Multi-Agent Growth Operating System"),
        ("Industry Domain & Focus:", "Indian D2C E-Commerce Logistics, Unit Economics & Cashflow Optimization"),
        ("Source Code Repository:", "https://github.com/jain44/GrowthOS.git (Branch: main)")
    ]
    for row_idx, (label, val) in enumerate(meta_data):
        c0 = meta_tbl.cell(row_idx, 0)
        c1 = meta_tbl.cell(row_idx, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, "F1F4F9")
        set_cell_background(c1, "FFFFFF")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(3)
        p0.paragraph_format.space_after = Pt(3)
        r0 = p0.add_run(label)
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(10.5)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(27, 54, 93)

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(3)
        r1 = p1.add_run(val)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = RGBColor(40, 40, 40)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(12)
    p_div.paragraph_format.space_after = Pt(6)

    # Overview
    add_heading_16(doc, "7.0 Architectural Overview & Module Organization")
    add_content_12(doc, "This annexure contains the complete production-grade source code for GrowthOS. The application is built using Next.js 16 (App Router & Turbopack), TypeScript 5, React 19, Tailwind CSS, PostgreSQL with Prisma ORM, NextAuth.js v5 (JWT authentication), and Groq Cloud Llama-3.3-70B-Versatile for autonomous multi-agent reasoning.", bold_prefix="Architecture Summary: ")
    add_content_12(doc, "The repository is engineered around a clean separation of concerns, dividing domain models, multi-agent AI heuristics, retrieval-augmented generation (RAG), authentication security, RESTful endpoints, and rich client-side dashboard visualizations into self-contained, typed modules.", bold_prefix="Engineering Design: ")

    # Directory Structure Block
    dir_structure = """GrowthOS/
├── prisma/
│   ├── schema.prisma                  # PostgreSQL database schema & relational entity models
│   └── seed.ts                        # Indian D2C 'Rang Studio' seed dataset & benchmarks
├── src/
│   ├── auth.ts                        # NextAuth.js v5 JWT auth provider & demo bypass
│   ├── agents/
│   │   └── shared.ts                  # Multi-agent contracts, schemas, and prioritization types
│   ├── lib/
│   │   ├── db.ts                      # Resilient Prisma database client & zero-latency fallback store
│   │   ├── utils.ts                   # Currency formatters (₹ Lakhs/Crores), health-score color helpers
│   │   └── ai/
│   │       ├── analyze.ts             # 7-Agent parallel diagnosis & executive synthesis engine
│   │       ├── simulate.ts            # Monte Carlo scenario simulator & financial forecasting
│   │       ├── report.ts              # Autonomous Executive CEO Briefing generator
│   │       └── rag.ts                 # Domain RAG knowledge retriever (Shiprocket, Razorpay benchmarks)
│   ├── app/
│   │   ├── (auth)/login/page.tsx      # Dark-mode enterprise login portal
│   │   ├── (dashboard)/
│   │   │   ├── dashboard/page.tsx     # Executive Dashboard (KPI metrics, margin cards, health score)
│   │   │   ├── dashboard/diagnose/page.tsx # Live 7-Agent Diagnostic runner & interactive playbook
│   │   │   ├── dashboard/simulate/page.tsx # What-If Strategy Simulator interface
│   │   │   ├── dashboard/report/page.tsx   # Strategic Board Report view & PDF exporter
│   │   │   └── dashboard/settings/page.tsx # Store settings & Groq engine parameters
│   │   └── api/
│   │       ├── analyze/route.ts       # POST /api/analyze (Trigger 7-agent diagnosis)
│   │       ├── simulate/route.ts      # POST /api/simulate (Run what-if scenario)
│   │       ├── report/route.ts        # POST /api/report (Generate CEO briefing)
│   │       └── business/route.ts      # GET/POST /api/business (Store profile CRUD)
│   └── components/layout/
│       └── dashboard-sidebar.tsx      # Sidebar navigation & same-origin logout handler
├── package.json                       # Core dependencies & npm scripts
└── next.config.ts                     # Turbopack & Next.js production build configuration"""
    add_content_heading_14(doc, "7.0.1 Repository File Structure & Core Components")
    add_code_block(doc, dir_structure)

    # 7.1 Schema
    add_heading_16(doc, "7.1 Database Architecture & Schema Definition (prisma/schema.prisma)")
    add_content_12(doc, "The relational schema models the entire lifecycle of an Indian D2C enterprise. It includes native attributes for Indian payment logistics including codRate (Cash on Delivery percentage), rtoRate (Return to Origin percentage), and prepaidPercentage, alongside multi-agent diagnosis runs, scenario simulations, board reports, historical financial snapshots, and RAG knowledge chunks.", bold_prefix="Schema Design: ")
    add_code_block(doc, read_file("prisma/schema.prisma"))

    # 7.2 Auth
    add_heading_16(doc, "7.2 Authentication & Session Security Engine (src/auth.ts)")
    add_content_12(doc, "GrowthOS implements NextAuth.js v5 (beta) utilizing an encrypted JSON Web Token (JWT) session strategy. To ensure absolute reliability during academic evaluations and offline presentations, the authorization engine features instant demo credential verification for demo@growthos.ai and safe relative-origin redirect handling.", bold_prefix="Security Implementation: ")
    add_code_block(doc, read_file("src/auth.ts"))

    # 7.3 Resilient DB
    add_heading_16(doc, "7.3 Resilient Database Client & Zero-Latency Hybrid Fallback Layer (src/lib/db.ts)")
    add_content_12(doc, "To guard against cloud pooler timeouts, network disconnections, or database dormancy during live evaluations, db.ts provides a resilient hybrid execution wrapper. When the primary PostgreSQL instance is active, queries execute through Prisma. If an infrastructure interruption or connection timeout occurs, the client automatically and transparently falls back to an in-memory repository containing the complete 'Rang Studio' Indian D2C dataset with zero latency.", bold_prefix="High-Availability Wrapper: ")
    add_code_block(doc, read_file("src/lib/db.ts"))

    # 7.4 Multi-Agent Diagnostic Engine
    add_heading_16(doc, "7.4 Multi-Agent AI Diagnostic Engine (src/lib/ai/analyze.ts)")
    add_content_12(doc, "This module coordinates the 7 domain-specific AI agents. Each agent (Finance, Operations, Marketing, Technology, Unit Economics, Inventory, Retention) receives specialized system prompts, benchmark retrieval contexts via RAG, and store parameters. The agents execute concurrently via Promise.allSettled() over the Groq Llama-3.3-70B inference engine, culminating in an executive synthesis pass that outputs prioritized problems, insights, and copy-pasteable operational playbooks.", bold_prefix="Multi-Agent Orchestration: ")
    add_code_block(doc, read_file("src/lib/ai/analyze.ts"))

    # 7.5 Shared Contracts
    add_heading_16(doc, "7.5 Multi-Agent Interface Contracts & Data Transfer Objects (src/agents/shared.ts)")
    add_content_12(doc, "Provides strict TypeScript interfaces for agent context inputs, diagnosis outputs, action items, operational playbook click paths, and executive synthesis scores.", bold_prefix="Type Safety Contracts: ")
    add_code_block(doc, read_file("src/agents/shared.ts"))

    # 7.6 Strategy Simulation Engine
    add_heading_16(doc, "7.6 Strategy Simulation & Counterfactual Financial Engine (src/lib/ai/simulate.ts)")
    add_content_12(doc, "The Strategy Simulator enables founders to test counterfactual scenarios (e.g., 'What happens if we reduce COD fee and mandate WhatsApp verification?'). It projects short-term margin effects, 12-month annualized financial impacts, and risk-adjusted confidence scores.", bold_prefix="Predictive Simulation: ")
    add_code_block(doc, read_file("src/lib/ai/simulate.ts"))

    # 7.7 Executive CEO Report Engine
    add_heading_16(doc, "7.7 Executive CEO Report Generation Engine (src/lib/ai/report.ts)")
    add_content_12(doc, "Synthesizes multi-agent diagnostic findings into an institutional board briefing. Generates executive summaries, identifies the single greatest operational hazard, highlights the highest-leverage growth unlock, and produces a phased 30-day tactical roadmap.", bold_prefix="Autonomous Briefing Generator: ")
    add_code_block(doc, read_file("src/lib/ai/report.ts"))

    # 7.8 RAG Retrieval Engine
    add_heading_16(doc, "7.8 Retrieval-Augmented Generation (RAG) Engine (src/lib/ai/rag.ts)")
    add_content_12(doc, "Retrieves verified Indian e-commerce industry benchmarks from curated logistics reports (Shiprocket 2024), checkout studies (Razorpay Magic Checkout), and fashion margin playbooks (Bain & Co) to ground all LLM recommendations in real market realities.", bold_prefix="Domain Grounding: ")
    add_code_block(doc, read_file("src/lib/ai/rag.ts"))

    # 7.9 API Endpoints
    add_heading_16(doc, "7.9 Backend REST API Endpoints & Server Action Handlers")
    
    add_content_heading_14(doc, "7.9.1 Diagnosis API Handler (src/app/api/analyze/route.ts)")
    add_content_12(doc, "Secured POST endpoint that validates user sessions, retrieves store financial metrics and historical snapshots, triggers the 7-agent diagnostic pipeline, and persists results.", bold_prefix="Endpoint Specification: ")
    add_code_block(doc, read_file("src/app/api/analyze/route.ts"))

    add_content_heading_14(doc, "7.9.2 Strategy Simulation API Handler (src/app/api/simulate/route.ts)")
    add_content_12(doc, "Executes counterfactual what-if strategic simulations against live brand financial profiles.", bold_prefix="Endpoint Specification: ")
    add_code_block(doc, read_file("src/app/api/simulate/route.ts"))

    add_content_heading_14(doc, "7.9.3 Executive Report API Handler (src/app/api/report/route.ts)")
    add_content_12(doc, "Generates and persists executive CEO summaries and structured action plans.", bold_prefix="Endpoint Specification: ")
    add_code_block(doc, read_file("src/app/api/report/route.ts"))

    add_content_heading_14(doc, "7.9.4 Store Business Profile & Snapshot API Handler (src/app/api/business/route.ts)")
    add_content_12(doc, "Provides CRUD operations for store onboarding data, operational logistics metrics, and monthly financial snapshots.", bold_prefix="Endpoint Specification: ")
    add_code_block(doc, read_file("src/app/api/business/route.ts"))

    # 7.10 Frontend UI
    add_heading_16(doc, "7.10 Frontend Executive Interface & Real-Time Visualization")

    add_content_heading_14(doc, "7.10.1 Executive Dashboard Page (src/app/(dashboard)/dashboard/page.tsx)")
    add_content_12(doc, "Renders key financial KPIs (Monthly Revenue, Net Expenses, Gross Profit, Operating Margin), store health score badge with color-coded alerts, and quick actions to launch diagnostics or simulations.", bold_prefix="Dashboard Interface: ")
    add_code_block(doc, read_file("src/app/(dashboard)/dashboard/page.tsx"))

    add_content_heading_14(doc, "7.10.2 Live 7-Agent Diagnostic Runner & Interactive Playbook (src/app/(dashboard)/dashboard/diagnose/page.tsx)")
    add_content_12(doc, "Interactive interface providing live agent status animations, priority-ranked friction point cards, and expandable software click paths with one-click copy templates.", bold_prefix="Diagnostic Interface: ")
    add_code_block(doc, read_file("src/app/(dashboard)/dashboard/diagnose/page.tsx"))

    add_content_heading_14(doc, "7.10.3 Sidebar Navigation & Same-Origin Session Handler (src/components/layout/dashboard-sidebar.tsx)")
    add_content_12(doc, "Sidebar navigation panel providing quick access to Overview, Diagnostic Suite, Strategy Simulator, CEO Reports, History, and Settings, paired with safe client-side same-origin logout.", bold_prefix="Navigation & Layout: ")
    add_code_block(doc, read_file("src/components/layout/dashboard-sidebar.tsx"))

    # Section 8: PPT Summary & Presentation Reference
    add_heading_16(doc, "8.0 Presentation Slide Deck & Demonstration Alignment")
    add_content_12(doc, "The academic PowerPoint presentation accompanying this technical report is titled 'GrowthOS_Presentation.pptx'. It comprises 11 structured slides aligned with university evaluation criteria, presenting the live demonstration narrative, multi-agent architecture diagrams, unit economics case study ('Rang Studio'), and full stack implementation highlights.", bold_prefix="Presentation Context: ")

    ppt_table = doc.add_table(rows=12, cols=3)
    ppt_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    ppt_headers = ["Slide #", "Slide Title", "Technical & Architectural Content Demonstrated"]
    for col_idx, h_text in enumerate(ppt_headers):
        cell = ppt_table.cell(0, col_idx)
        set_cell_background(cell, "1B365D")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(h_text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    slides_info = [
        ("Slide 1", "Title Slide: GrowthOS", "Autonomous Multi-Agent Growth Operating System for Indian E-Commerce Brands."),
        ("Slide 2", "The Indian D2C Crisis", "High COD dependency (60-70%), crushing RTO freight loss (20-30%), rising Meta CAC, and low contribution margins."),
        ("Slide 3", "Introducing GrowthOS", "Real-time AI board of directors diagnosing, simulating, and optimizing unit economics autonomously."),
        ("Slide 4", "Multi-Agent System Architecture", "Parallel execution of 7 specialized domain agents via Groq Cloud Llama-3.3-70B & RAG domain grounding."),
        ("Slide 5", "The 7 Autonomous Growth Agents", "Finance, Operations, Marketing, Technology, Unit Economics, Inventory/Supply Chain, and Retention/LTV."),
        ("Slide 6", "Full-Stack Technology Stack", "Next.js 16 (Turbopack), TypeScript 5, Tailwind CSS, PostgreSQL, Prisma ORM, NextAuth.js v5 JWT."),
        ("Slide 7", "Case Study: Rang Studio", "₹18.4L MRR ethnic wear brand; 62% COD, 22% RTO. Health score 62/100 analyzed across 6 months of historical data."),
        ("Slide 8", "What-If Strategy Simulator", "Counterfactual scenario testing: +₹1.4L/mo profit recovery by shifting 25% COD to prepaid with WhatsApp NDR."),
        ("Slide 9", "Autonomous CEO Briefing & PDF", "One-click institutional board reports with prioritized operational playbooks, software click paths, and copy templates."),
        ("Slide 10", "Business Impact & Metrics", "15% COD fee shifted 15% orders to prepaid; 6% RTO reduction; ₹1.8L monthly cash flow unlocked."),
        ("Slide 11", "Future Roadmap & Conclusion", "Automated Shopify/Shiprocket API webhooks, fine-tuned Indian commerce SLMs, and multi-tenant enterprise orchestration.")
    ]

    for row_idx, (s_num, s_title, s_desc) in enumerate(slides_info, start=1):
        c0 = ppt_table.cell(row_idx, 0)
        c1 = ppt_table.cell(row_idx, 1)
        c2 = ppt_table.cell(row_idx, 2)
        c0.width = Inches(1.0)
        c1.width = Inches(2.2)
        c2.width = Inches(3.3)
        
        bg = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        set_cell_background(c0, bg)
        set_cell_background(c1, bg)
        set_cell_background(c2, bg)

        for c, text, is_bold in [(c0, s_num, True), (c1, s_title, True), (c2, s_desc, False)]:
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
            r.font.bold = is_bold
            r.font.color.rgb = RGBColor(30, 30, 30)

    # Save to submissions and project root
    out_path_sub = os.path.join(project_root, "submissions", "FSD1_ANNEXURE_SOURCE_CODE.docx")
    out_path_root = os.path.join(project_root, "FSD1_ANNEXURE_SOURCE_CODE.docx")
    
    os.makedirs(os.path.dirname(out_path_sub), exist_ok=True)
    doc.save(out_path_sub)
    doc.save(out_path_root)
    print(f"SUCCESS: Generated Word Annexure Document:\n- {out_path_sub}\n- {out_path_root}")

build_annexure_doc()
